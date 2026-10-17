# -*- coding: UTF-8 -*-
#
# Reproducer for CVE-2024-53924 / GHSA-pw67-xjhq-389w: a crafted formula in an
# untrusted workbook could execute arbitrary Python when pycel evaluated it,
# because an Excel function name with no pycel implementation was emitted as a
# raw Python call and run in a namespace that still had the full Python builtins
# (eval, exec, __import__, open, ...).
#
# The payload here is harmless: it only tries to create a marker file inside the
# test's own temp directory. It runs no shell command and touches no network.
# On the unpatched code the marker appears (test FAILS); on the fix the builtin
# name is undefined, pycel reports it as an unknown function, and no marker is
# written (test PASSES).

import openpyxl
import pytest

from pycel import ExcelCompiler


@pytest.mark.parametrize('builtin', ['eval', 'exec'])
def test_crafted_formula_cannot_execute_python(tmp_path, builtin):
    marker = tmp_path / 'pwned-{}.txt'.format(builtin)
    # Harmless body: create a marker file in this test's temp dir. Nothing else.
    # Forward slashes: pathlib accepts them on Windows and they avoid
    # backslash escape sequences inside the emitted Python string literal.
    body = ("__import__('pathlib').Path('{}').write_text('x')"
            .format(str(marker).replace('\\', '/')))
    payload = '{}("{}")'.format(builtin, body.replace('"', '\\"'))

    xlsx = tmp_path / 'evil-{}.xlsx'.format(builtin)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws['A1'] = 200
    # The two shapes the advisory and its gist call out: inside IF, and *.
    ws['A2'] = '=IF(A1=200, {}, A1)'.format(payload)
    ws['A3'] = '=A1*{}'.format(payload)
    wb.save(str(xlsx))

    excel = ExcelCompiler(str(xlsx))
    for cell in ('{}!A2'.format(ws.title), '{}!A3'.format(ws.title)):
        # The fix makes the builtin name undefined, so pycel raises (unknown
        # function); the unpatched code runs the payload and returns a value.
        # Either way we only care that the payload did not run, so swallow any
        # evaluation error and assert on the marker.
        try:
            excel.evaluate(cell)
        except Exception:
            pass

    assert not marker.exists(), (
        'crafted formula executed Python via {!r}'.format(builtin))


def test_benign_formula_still_evaluates(tmp_path):
    # The fix must not break ordinary formulas.
    xlsx = tmp_path / 'ok.xlsx'
    wb = openpyxl.Workbook()
    ws = wb.active
    ws['A1'] = 200
    ws['A2'] = '=IF(A1=200, A1*2, A1)'
    ws['A3'] = '=SUM(A1:A2)'
    wb.save(str(xlsx))
    excel = ExcelCompiler(str(xlsx))
    assert excel.evaluate('{}!A2'.format(ws.title)) == 400
    assert excel.evaluate('{}!A3'.format(ws.title)) == 600
